import re
import sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_run_font(run, name_cn="宋体", size=11, bold=None, color=None):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name_cn)
    if bold is not None:
        run.font.bold = bold
    if color:
        run.font.color.rgb = color


def add_rich_text(paragraph, text, size=11, cn_font="宋体", color=None):
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, cn_font, size, bold=True, color=color)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, cn_font, size, bold=False, color=color)


def shade(cell, hexcolor="F1EFE8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tc_pr.append(shd)


def build(md_path, docx_path):
    lines = open(md_path, encoding="utf-8").read().split("\n")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(4)

    i = 0
    first_h1 = True
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            if level == 1 and first_h1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(6)
                add_rich_text(p, text, size=20, cn_font="黑体")
                first_h1 = False
            elif level == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_before = Pt(16)
                p.paragraph_format.space_after = Pt(8)
                add_rich_text(p, text, size=16, cn_font="黑体")
            elif level == 2:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(6)
                add_rich_text(p, text, size=14, cn_font="黑体")
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(4)
                add_rich_text(p, text, size=12, cn_font="黑体")
            i += 1
            continue

        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for tl in table_lines:
                cells = [c.strip() for c in tl.strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    continue
                rows.append(cells)
            if not rows:
                continue
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_idx, row in enumerate(rows):
                for c_idx in range(cols):
                    cell = table.cell(r_idx, c_idx)
                    text = row[c_idx] if c_idx < len(row) else ""
                    para = cell.paragraphs[0]
                    para.paragraph_format.line_spacing = 1.2
                    para.paragraph_format.space_after = Pt(1)
                    add_rich_text(para, text, size=9.5, color=RGBColor(0x2C, 0x2C, 0x2A))
                    if r_idx == 0:
                        shade(cell)
                        for run in para.runs:
                            run.font.bold = True
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_rich_text(p, stripped[2:])
            i += 1
            continue

        if stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(2)
            add_rich_text(p, stripped[2:], color=RGBColor(0x5F, 0x5E, 0x5A))
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(2)
            add_rich_text(p, m.group(2))
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        add_rich_text(p, stripped)
        i += 1

    doc.save(docx_path)
    print("saved:", docx_path)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
